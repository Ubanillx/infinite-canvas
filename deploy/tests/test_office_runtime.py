"""Run only against the real final Docker runtime, including its copied tools."""
import importlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from PIL import Image


class OfficeRuntimeTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(prefix="canvas-office-")
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)

    def inspect(self, name):
        return subprocess.run(["canvas-image-info", str(self.root / name)], capture_output=True, text=True)

    def test_expected_nonroot_environment_and_office_dependencies(self):
        import os
        self.assertNotEqual(os.getuid(), 0)
        self.assertEqual(Path(sys.prefix), Path("/opt/office"))
        for module in ["PIL", "docx", "openpyxl", "pptx", "pypdf", "reportlab", "pandas", "requests"]:
            with self.subTest(module=module):
                importlib.import_module(module)
        for command in ["python", "python3", "canvas-image-info", "pdftotext", "pdftoppm", "pdfinfo", "libreoffice", "pandoc", "tesseract", "ffmpeg", "ffprobe", "curl", "jq", "rg", "file", "zip", "unzip"]:
            self.assertIsNotNone(shutil.which(command), command)
        fonts = subprocess.check_output(["fc-list", ":lang=zh", "family"], text=True)
        self.assertIn("Noto", fonts)
        self.assertIn("chi_sim", subprocess.check_output(["tesseract", "--list-langs"], text=True))

    def test_real_alpha_and_palette_transparency_without_modifying_originals(self):
        Image.new("RGBA", (2, 2), (255, 255, 255, 255)).save(self.root / "opaque.png")
        transparent = Image.new("RGBA", (2, 2), (0, 0, 0, 0))
        transparent.putpixel((1, 1), (255, 0, 0, 128))
        transparent.save(self.root / "alpha.png")
        palette = Image.new("P", (2, 2), 0)
        palette.save(self.root / "palette.png", transparency=0)
        for name, expected in [("opaque.png", False), ("alpha.png", True), ("palette.png", True)]:
            with self.subTest(name=name):
                before = (self.root / name).read_bytes()
                result = self.inspect(name)
                self.assertEqual(result.returncode, 0, result.stderr)
                info = json.loads(result.stdout)
                self.assertEqual(info["has_transparent_pixels"], expected)
                self.assertEqual(info["width"], 2)
                self.assertEqual((self.root / name).read_bytes(), before)
        alpha = json.loads(self.inspect("alpha.png").stdout)
        self.assertEqual(alpha["transparent_pixels"], 3)
        self.assertEqual(alpha["partial_alpha_pixels"], 1)

    def test_json_named_png_fails_with_nonzero_exit(self):
        (self.root / "error.png").write_text('{"error":"403 Forbidden"}')
        result = self.inspect("error.png")
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(result.stdout, "")
        self.assertIn("error", json.loads(result.stderr))

    def test_office_files_and_headless_pdf_conversion(self):
        from docx import Document
        from openpyxl import Workbook, load_workbook
        from pptx import Presentation
        from pypdf import PdfReader
        document = Document()
        document.add_paragraph("Canvas office runtime")
        source = self.root / "sample.docx"
        document.save(source)
        before = source.read_bytes()
        workbook = Workbook()
        workbook.active["A1"] = "Canvas"
        workbook.save(self.root / "sample.xlsx")
        self.assertEqual(load_workbook(self.root / "sample.xlsx").active["A1"].value, "Canvas")
        slides = Presentation()
        slides.slides.add_slide(slides.slide_layouts[0])
        slides.save(self.root / "sample.pptx")
        self.assertEqual(len(Presentation(self.root / "sample.pptx").slides), 1)
        result = subprocess.run(["libreoffice", f"-env:UserInstallation={(self.root / 'profile').as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(self.root), str(source)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Canvas office runtime", PdfReader(self.root / "sample.pdf").pages[0].extract_text())
        self.assertEqual(source.read_bytes(), before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
